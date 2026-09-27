"""Evidence-linked, four-page Chinese editorial manuscript with native math."""
from layout import *
from docx.enum.text import WD_BREAK

def column_break(doc):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1);p.add_run().add_break(WD_BREAK.COLUMN)

def algorithm_current(doc):
    p=heading(doc,'算法 1  后续相关披露规划')
    b=OxmlElement('w:pBdr');n=OxmlElement('w:top');n.set(qn('w:val'),'single');n.set(qn('w:sz'),'6');b.append(n);p._p.get_or_add_pPr().append(b)
    lines=['Input: known choices A, full history H, field budget B',
      'Output: optimal registered plan, or explicit failure',
      '1:  G ← ComponentsSharingNewTokens(A, H)',
      '2:  for each component g in G do',
      '3:      F[g] ← ∅; U ← KnownSuffixTokens(g, H)',
      '4:      for each attainable field bound b ≤ B do',
      '5:          L ← {(live=∅, retired=0, work=0, fields=0)}',
      '6:          for each step i in original order of g do',
      '7:              N ← ∅',
      '8:              for each label l in L and choice a in A[i] do',
      '9:                  if l.fields + a.fields > b then continue',
      '10:                 X ← l.live ∪ (Tokens(a) ∖ H)',
      '11:                 live′ ← X ∩ U[i+1]',
      '12:                 retired′ ← l.retired + |X ∖ U[i+1]|',
      '13:                 N ← N ∪ Extend(l, a, live′, retired′)',
      '14:             L ← PruneWithinEqualLiveSets(N)',
      '15:             if |L| > LIMIT then return LIMIT_REACHED',
      '16:         if L ≠ ∅ then AddBest(F[g], L)',
      '17:     F[g] ← PruneCostFrontier(F[g])',
      '18: P ← CombineComponentsUnderOneBudget(F, B)',
      '19: if P is empty then return NO_FEASIBLE_PLAN',
      '20: return RestoreOriginalStepOrder(Best(P))']
    for i,line in enumerate(lines):
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05;p.paragraph_format.keep_with_next=i<len(lines)-1
        run(p,line,8,bold=i<2)
        if i==len(lines)-1:
            b=OxmlElement('w:pBdr');n=OxmlElement('w:bottom');n.set(qn('w:val'),'single');n.set(qn('w:sz'),'6');b.append(n);p._p.get_or_add_pPr().append(b)

def build():
    old=json.loads((ROOT/'evidence.json').read_text(encoding='utf8'))['financial']
    study=json.loads((ROOT/'../../evidence/frontier-study/summary.json').read_text(encoding='utf8'))
    final=study['final'];m=final['methods'];integration=study['integration']
    doc=Document();page(doc.sections[0],1)
    doc.core_properties.title='面向动态智能体任务的按步执行与信息共享';doc.core_properties.author='Yaojunzhe Zhai; Chao-Hsien Hsieh';doc.core_properties.comments=''
    normal=doc.styles['Normal'];normal.font.name='Linux Libertine';normal.font.size=Pt(9);normal.paragraph_format.space_after=Pt(0)
    p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3)
    b=OxmlElement('w:pBdr');n=OxmlElement('w:bottom');n.set(qn('w:val'),'nil');b.append(n);p._p.get_or_add_pPr().append(b)
    run(p,'面向动态智能体任务的按步执行与信息共享',14,True,cjk='SimHei')
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(2);run(p,'Yaojunzhe Zhai¹  ·  Chao-Hsien Hsieh²',9)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4);run(p,'¹ Monash University Malaysia   ² Xi’an International University',8)
    col_section(doc)
    heading(doc,'摘要')
    paragraph(doc,'智能体访问本地资料、云端模型和外部工具时，后续步骤及其所需输入可能随执行而改变。逐次选择发送量最少的视图，未必使整段任务的累计披露最少；直接枚举视图组合又容易使规划状态膨胀。本文介绍一个按步执行原型：将执行位置与接收方视图作为共同候选，按来源版本核验请求，并在传输预算内选择当前已知步骤的后续方案。规划器仅保留后续仍可能使用的披露项，将其余项转为累计成本，再分解不共享新增披露项的步骤。48 个构造负载的 144 个预算设置均与独立整数规划得到相同最优值；已有真实模型计划的重新执行验证了控制器与 HTTP 接收记录的一致性。双栏界面允许观众检查每一步的选择、实际发送内容及状态变化后的重算。',first=False)
    paragraph(doc,'关键词：智能体执行；接收方视图；披露状态压缩；预算规划；系统演示。',first=False,size=8)
    heading(doc,'1  问题与系统目标')
    paragraph(doc,'企业智能体可能先检查内部记录，再向云端解释模型或业务工具发出请求。不同接收方需要不同输入：模型解释可以使用业务事实，检索工具只需查询条件，数值服务则可接收原值或已在本地计算的中间值。固定发送完整上下文会扩大共享范围，而固定裁剪又可能遗漏后来出现的需求。系统必须在每个已知步骤重新选择执行方式和输入。')
    paragraph(doc,'选择还受历史影响。两步均可发送同一对原值，或分别发送不同派生值；逐步贪心可能先选一个派生值，随后再披露两个新值。连续使用原值方案只增加两项不同披露，却可能产生更多重复传输。因此，累计不同披露项与传输字段次数需要分开控制。状态或能力变化后，这个选择还必须与结果缓存保持一致。')
    paragraph(doc,'PlanTwin [1]、PrivScope [2] 和 MINIM [3] 已研究规划抽象、任务范围与必要性裁剪。本文聚焦登记视图之间的有限选择：如何在不删除实际披露历史的前提下压缩搜索状态，并把选择落实为可核查的执行。状态压缩借鉴前沿搜索 [4]，不将动态规划或本地门控本身作为首次提出的技术。')
    paragraph(doc,'本文提供三项可检查的结果：面向接收方并集成本和传输预算的精确规划；与版本核验、局部重算和真实 HTTP 请求衔接的控制器；分开报告优化正确性、规划开销和模型任务结果的开放演示。有限候选由应用登记，候选的语义等价性不由规划器自行推断。')
    col_section(doc,count=1)
    full_figure(doc,'assets/system.png','图 1  按步执行的系统结构。原始资料保留在本地；每个新步骤或状态变化触发重新判定，外部执行通过登记出口完成，结果返回本地。')

    col_section(doc,WD_SECTION_START.NEW_PAGE,2)
    heading(doc,'2  后续相关披露规划')
    heading(doc,'2.1  有限候选与优化目标',2)
    paragraph(doc,'状态 S 保存来源版本、授权、能力和有效计算结果。当前已知的 m 个步骤各有有限候选集 Aᵢ；候选指定执行位置、接收方、输入视图、工作量与字段传输量。原值或派生值的标识包含接收方、表达式及来源版本。令 H 为此前实际发送的完整历史，E(a) 为候选的披露项，W、T 分别为总工作量和字段次数，则：')
    math(doc,'<munder><mtext>min lex</mtext><mrow><mi>p</mi><mo>∈</mo><mi>P</mi><mo>,</mo><mi>T</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>≤</mo><mi>B</mi></mrow></munder><mo>(</mo><mo>|</mo><munder><mo>⋃</mo><mrow><mi>a</mi><mo>∈</mo><mi>p</mi></mrow></munder><mi>E</mi><mo>(</mo><mi>a</mi><mo>)</mo><mo>∖</mo><mi>H</mi><mo>|</mo><mo>,</mo><mi>W</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>,</mo><mi>T</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>)</mo>',1)
    paragraph(doc,'P 仅含授权、能力与输入需求均满足的方案。B 是同阶段预算；重复发送仍消耗字段次数，历史披露不因重算而撤销。本地完成不增加披露项。该成本衡量显式共享项目，不衡量接收方可能推断出的全部信息。')
    heading(doc,'2.2  保留仍与后续有关的状态',2)
    paragraph(doc,'设 Eᵢ 为当前前缀的新增披露集合，Uᵢ 为已知后缀所有候选可能披露的、尚不在 H 中的项目。完整集合保存已不可能复用的项目，会使本应等价的前缀无法合并。我们保留活动集合 Lᵢ，并用 zᵢ 计入退出该集合的项目：')
    math(doc,'<msub><mi>L</mi><mi>i</mi></msub><mo>=</mo><msub><mi>E</mi><mi>i</mi></msub><mo>∩</mo><msub><mi>U</mi><mi>i</mi></msub><mo>,</mo><mspace width="0.4em"/><msub><mi>z</mi><mi>i</mi></msub><mo>=</mo><mo>|</mo><msub><mi>E</mi><mi>i</mi></msub><mo>∖</mo><msub><mi>U</mi><mi>i</mi></msub><mo>|</mo>',2)
    paragraph(doc,'对任意共同后缀的新增披露集合 F⊆Uᵢ，退出的项目与后缀不相交，因此有：')
    math(doc,'<mo>|</mo><msub><mi>E</mi><mi>i</mi></msub><mo>∪</mo><mi>F</mi><mo>|</mo><mo>=</mo><msub><mi>z</mi><mi>i</mi></msub><mo>+</mo><mo>|</mo><msub><mi>L</mi><mi>i</mi></msub><mo>∪</mo><mi>F</mi><mo>|</mo>',3)
    paragraph(doc,'加入候选后令 X=Lᵢ∪(E(a)∖H)，按下一边界更新。项目在最后一次可能出现后退出，之后不会被再次计数：')
    math(doc,'<msup><mi>L</mi><mo>′</mo></msup><mo>=</mo><mi>X</mi><mo>∩</mo><msub><mi>U</mi><mrow><mi>i</mi><mo>+</mo><mn>1</mn></mrow></msub><mo>,</mo><mspace width="0.3em"/><msup><mi>z</mi><mo>′</mo></msup><mo>=</mo><msub><mi>z</mi><mi>i</mi></msub><mo>+</mo><mo>|</mo><mi>X</mi><mo>∖</mo><msub><mi>U</mi><mrow><mi>i</mi><mo>+</mo><mn>1</mn></mrow></msub><mo>|</mo>',4)
    paragraph(doc,'标签为 (L,z,w,b)。相同 L 下，若一个标签传输量不高，且已累计的 (z,w) 按字典序更优，则另一标签可被删除。严格较少的 z 优先于任何工作量差异：')
    math(doc,'<msub><mi>b</mi><mn>1</mn></msub><mo>≤</mo><msub><mi>b</mi><mn>2</mn></msub><mo>∧</mo><mo>(</mo><msub><mi>z</mi><mn>1</mn></msub><mo>,</mo><msub><mi>w</mi><mn>1</mn></msub><mo>)</mo><msub><mo>≤</mo><mtext>lex</mtext></msub><mo>(</mo><msub><mi>z</mi><mn>2</mn></msub><mo>,</mo><msub><mi>w</mi><mn>2</mn></msub><mo>)</mo>',5)
    paragraph(doc,'式（3）保证追加共同后缀后该关系仍成立。若后续可行性不依赖前缀所选表示，剪枝保留至少一个最优解。设跨边界可能复用的项目数最多为 ω，字段预算为整数 B，则保留标签数满足：')
    math(doc,'<mo>|</mo><mi>Labels</mi><mo>|</mo><mo>≤</mo><msup><mn>2</mn><mi>ω</mi></msup><mo>(</mo><mi>B</mi><mo>+</mo><mn>1</mn><mo>)</mo>',6)
    paragraph(doc,'这是依赖边界宽度和数值预算的界，不是一般多项式保证。实现还需要处理组内支配比较；超过 4,096 个标签会明确报告资源上限。')
    column_break(doc)
    heading(doc,'2.3  分解披露关联，保持执行顺序',2)
    paragraph(doc,'两个步骤若有候选共享新增披露项，便在成本图中相连。不同连通分量的新增集合不相交，披露量、工作量和传输量可以相加。对每个分量求不同字段上限下的最优候选，再在统一 B 下组合成本前沿。任何全局最优解都可用相同字段上限下的分量最优解替换，因此分解保持最优值。')
    paragraph(doc,'这里分解的是成本，不是执行过程。返回方案恢复原始步骤顺序，数据依赖和实际工具调用顺序保持不变；共享副作用或输出不等价的候选不适用该分解。完整证明和独立整数规划模型随代码提供。')
    algorithm_current(doc)
    heading(doc,'2.4  从方案到实际请求',2)
    paragraph(doc,'控制器逐步执行所选方案。发送前重新核对授权、能力和依赖来源版本，序列化后记录请求摘要；实际 HTTP 接收端保存正文与结果。新执行器在发起发送时记入披露历史，即使响应失败也不删除可能已经外发的信息。摘要用于核对本地记录，不是云服务商的独立证明。')
    paragraph(doc,'源值更新只使依赖该源的结果失效；能力变化使候选重新选择，但不自动删除仍有效的结果。新步骤出现时，以完整 H 和当前已知后缀重新规划，曾退出搜索状态的项目仍在 H 中。系统不假定能够看见未来尚未出现的步骤。每个稳定阶段共用一个预算余额，不能在每次重规划时重新领取额度。')

    col_section(doc,WD_SECTION_START.NEW_PAGE,1)
    full_figure(doc,'../../evidence/frontier-showcase/paper-ui.png','图 2  新执行器的实际界面。① 从保存的真实模型响应创建步骤；② 本地能力变化后保留有效结果并重规划；③ 核对所选输入、发送内容及 HTTP 接收记录。截图裁去导航并添加编号，不增加实验样本数。',width=6.65)
    col_section(doc)
    heading(doc,'3  观众如何体验系统')
    paragraph(doc,'界面沿用客户端的双栏任务布局：左侧自动推进执行，右侧显示当前步骤的接收方、输入选择、本地保留字段及请求正文。观众可以查看旧步骤并恢复跟随；整个任务无需逐步点击确认。中文与英文界面使用同一执行记录，展示时间与实际计算耗时分开。')
    paragraph(doc,'首先运行“真实模型计划·本地能力变化”。保存的 DeepSeek 调用先取得表结构，再提出相加、相加、相除的程序。首步在本地完成后，登记事件撤销本地加减能力；系统保留首步结果，向 HTTP 算子发送剩余计算所需的输入。切换“条件变化后全部重跑”，可直接观察多出的计算和请求。')
    paragraph(doc,'随后比较已用与未用源值更新，检查哪些结果失效、哪些仍可复用。观众在同一页面切换压缩规划、完整状态规划和逐步贪心，查看所选输入与预算；这些短模型计划的最优值相同，页面不会为制造差异而篡改结果。较长任务的规模差异由公开规划实验单独展示。')
    paragraph(doc,'无密钥模式重新执行计算与 HTTP 请求，只复用已公开的模型计划；读者也可用自备密钥生成新计划。能力撤销和源值更新是预定的受控条件，不是模型发现的故障。算术服务为回环 HTTP 接收端，混合能力设置用于研究执行选择，不意味着基本算术必须外发。')
    column_break(doc)
    heading(doc,'4  实验设计与复核方式')
    paragraph(doc,'实验分为算法问题、真实模型接入和受控执行三层。最终算法在开发检查后冻结源代码、生成规则和工作负载摘要；使用新的固定种子产生 48 个负载，每个分别设置不增加、增加 25% 和增加 100% 的字段预算，共 144 个设置。三档预算共享任务输入，不作为独立自然任务计数。')
    paragraph(doc,'六类负载覆盖独立选择、局部共享、远距离重用、链式重叠、不同接收方和部分本地能力，长度为 8、16、32、64 步。工作量与已有历史重新抽样。这些是构造的登记候选，不是生产工作流，也不能估计实际隐私泄露频率。开发阶段的失败和完整记录同样保留。')
    paragraph(doc,'最终对照为共享候选的逐步贪心、仅退出无关项的消融，以及 SciPy/HiGHS 独立整数规划 [6]；开发阶段另保留完整集合前沿的结果。整数模型对候选选择和披露项并集分别设二元变量，约束共用字段预算。比较最优目标三元组、资源上限、峰值状态和规划时间；所有失败计入分母。')
    paragraph(doc,f'接入实验复用此前冻结的 FinQA [5] 材料：{old["questions"]} 道公开题、{old["plans"]} 份真实模型计划及 {old["modelCalls"]} 次原始请求。新执行器与原执行器在五种条件下重新执行，共 {integration["runs"]} 次、{integration["httpRequests"]} 条 HTTP 接收记录，没有新增模型调用。另用 250 个小问题穷举验证规划最优值；它们是算法测试，不能并入模型正确率。')

    col_section(doc,WD_SECTION_START.NEW_PAGE,1)
    full_figure(doc,'../../evidence/validation-v8/tradeoffs.zh-CN.png','图 3  既有真实模型计划的受控执行结果，每种方法 240 次。结果复用减少 HTTP 调用与字段传输；该短任务批次中贪心与后续规划一致。此图验证执行层，表 1—2 单独评价新的规划器。',width=6.78)
    col_section(doc)
    heading(doc,'5  结果与分析')
    table(doc,'表 1  最终规划结果（峰值取已完成设置）',[
        ['方法','求解成功','同最优值','峰值状态'],
        ['逐步贪心','144/144',str(m['greedy']['optimalMatches'])+'/144','1'],
        ['仅退出无关项',str(m['live-frontier']['completed'])+'/144',str(m['live-frontier']['optimalMatches'])+'/144',str(m['live-frontier']['maxPeak'])],
        ['完整压缩机制','144/144','144/144',str(m['disclosure-frontier']['maxPeak'])],
        ['独立整数规划','144/144','144/144','—']], [1.08,.74,.77,.74])
    paragraph(doc,'完整机制在所有设置上与整数规划一致；仅退出无关项仍有 15 个设置达到标签上限。分量分解和按字典序支配删除恢复了这些可解设置。整数规划是独立最优值参照，表中的一致不等于本方法在所有任务上比通用求解器更快。',first=False)
    table(doc,'表 2  披露与传输的交换关系',[
        ['预算增幅','不同披露项','字段次数'],
        *[[f'{int(r["slack"]*100)}%',f'{r["disclosures"]} / {r["greedyDisclosures"]}',f'{r["fields"]} / {r["greedyFields"]}'] for r in final['budgets']]], [0.75,1.32,1.26])
    paragraph(doc,'每格为“完整机制 / 贪心”，各行含 48 个相同负载。不增加预算时不同披露项减少 1.9%；预算放宽 25% 时减少 14.6%，实际传输却增加 18.9%。减少不同项目不等于减少通信，更不等于同比减少隐私风险。',first=False,size=8)
    paragraph(doc,f'新执行器的 {integration["runs"]} 次运行均与给定程序结果一致，接收正文与发送摘要逐条相符。但原始模型答案与题目标签仅 {old["labelMatches"]}/{old["plans"]} 一致；规划压缩不会修复模型选错年份或公式。既有实验的 HTTP 调用从 {old["restartCalls"]} 次降到 {old["reuseCalls"]} 次来自结果复用，不能归为本轮规划算法的收益。')
    if 'serialTiming' in study:
        tm=study['serialTiming']['methods']
        paragraph(doc,f'串行复测每个设置五次，先取各设置中位数。完整机制与整数规划的总体中位数分别为 {tm["disclosure-frontier"]["medianMs"]:.2f} 和 {tm["milp"]["medianMs"]:.2f} 毫秒；最大设置中位数为 {tm["disclosure-frontier"]["maxMedianMs"]:.1f} 和 {tm["milp"]["maxMedianMs"]:.1f} 毫秒。硬件及逐次时间随记录提供。',size=8,first=False)
    column_break(doc)
    heading(doc,'6  适用范围与开放实现')
    paragraph(doc,'方法适用于输出等价、成本与后续可行性可登记的有限候选。后缀范围大、披露关联紧密或预算很大时仍可能遇到资源上限；当前不优化步骤顺序，也不保证未知未来的全局最优。登记依赖和版本维护必须可信，任意网络绕过不在本原型的保证范围内。')
    paragraph(doc,'项目开源独立研究页面、规划器、接收端、冻结请求、原始记录及核验脚本，商业客户端不包含在内。默认演示无需模型密钥；公开记录可离线重算，重新调用模型需读者自己的凭据。论文中的数字由记录生成，开发失败、人审材料和既有业务适配器另附，未将尚未完成的人审当成质量证据。')
    paragraph(doc,'源码与操作说明：',first=False,size=8,space=0)
    paragraph(doc,'https://github.com/Zane-0260907/stepwise-disclosure-demo',first=False,size=7.7,space=2)
    heading(doc,'参考文献')
    refs=[
      '[1] G. Yu 等. PlanTwin: Privacy-Preserving Planning Abstractions for Cloud-Assisted LLM Agents. arXiv:2603.18377, 2026.',
      '[2] S. R. Seeam 等. PrivScope: Task-scoped Disclosure Control for Hybrid Agentic Systems. arXiv:2605.16630, 2026.',
      '[3] H. Yu 等. MINIM: Privacy-Aware Minimal View for Agents via Trusted Local Sanitization. arXiv:2606.13949, 2026.',
      '[4] J. Kawahara 等. Frontier-Based Search for Enumerating All Constrained Subgraphs with Compressed Representation. IEICE Trans. E100-A(9), 1773–1784, 2017.',
      '[5] Z. Chen 等. FinQA: A Dataset of Numerical Reasoning over Financial Data. EMNLP, 3697–3711, 2021.',
      '[6] P. Virtanen 等. SciPy 1.0: Fundamental Algorithms for Scientific Computing in Python. Nature Methods 17, 261–272, 2020. 整数规划使用 SciPy 1.15.3 的 HiGHS 接口.'
    ]
    for text in refs:
        p=paragraph(doc,text,first=False,size=8,space=1,align=WD_ALIGN_PARAGRAPH.LEFT);p.paragraph_format.left_indent=Inches(.14);p.paragraph_format.first_line_indent=Inches(-.14)
    view=OxmlElement('w:view');view.set(qn('w:val'),'print');doc.settings.element.insert(0,view)
    doc.save(OUT);print(OUT)

if __name__=='__main__':build()
