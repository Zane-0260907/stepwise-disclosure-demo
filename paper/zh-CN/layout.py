"""Build the evidence-aligned Chinese SIGMOD demo editorial draft."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parent
OUT = ROOT / '按步执行与信息共享_中文最新稿.docx'
XSL = Path(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL')


def page(section, columns):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0417)
    section.bottom_margin = Inches(1.1111)
    section.left_margin = Inches(.75)
    section.right_margin = Inches(.75)
    section.header_distance = Inches(.3)
    section.footer_distance = Inches(.35)
    cols = section._sectPr.xpath('./w:cols')
    cols = cols[0] if cols else OxmlElement('w:cols')
    cols.set(qn('w:num'), str(columns))
    cols.set(qn('w:space'), '480')
    if not section._sectPr.xpath('./w:cols'):
        section._sectPr.append(cols)


def run(p, value, size=9, bold=False, italic=False, cjk='SimSun'):
    r = p.add_run(value)
    r.font.name = 'Linux Libertine'
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor(0, 0, 0)
    rpr = r._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    fonts.set(qn('w:eastAsia'), cjk)
    return r


def paragraph(doc, value, first=True, size=9, space=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    f = p.paragraph_format
    f.first_line_indent = Inches(.16 if first else 0)
    f.space_after = Pt(space)
    f.line_spacing = 1.025
    run(p, value, size)
    return p


def heading(doc, label, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6 if level == 1 else 4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run(p, label, 9.4 if level == 1 else 9, True, cjk='SimHei')
    return p


def full_figure(doc, path, caption, width=6.78):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ROOT / path), width=Inches(width))
    cp = doc.add_paragraph()
    cp.paragraph_format.space_after = Pt(3)
    cp.paragraph_format.keep_with_next = False
    run(cp, caption, 8)


def col_section(doc, start=WD_SECTION_START.CONTINUOUS, count=2):
    section = doc.add_section(start)
    page(section, count)
    return section


def equation(doc, mathml, number):
    cache_path=ROOT / 'equations.omml.json'
    cache=json.loads(cache_path.read_text(encoding='utf8')) if cache_path.exists() else {}
    if mathml not in cache:
        if not XSL.exists(): raise RuntimeError('Equation not in cache; update the native OMML cache when changing formulas.')
        transform=etree.XSLT(etree.parse(str(XSL)))
        omml=transform(etree.fromstring(mathml.encode('utf8')))
        cache[mathml]=etree.tostring(omml.getroot(),encoding='unicode')
        cache_path.write_text(json.dumps(cache,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p._p.append(etree.fromstring(cache[mathml].encode('utf8')))
    run(p, f'  ({number})', 8)


def math(doc, content, n):
    equation(doc, '<math xmlns="http://www.w3.org/1998/Math/MathML" display="block"><mrow>'+content+'</mrow></math>', n)


def algorithm(doc):
    p=heading(doc,'算法 1  按步执行与受限事实补充')
    borders=OxmlElement('w:pBdr');top=OxmlElement('w:top');top.set(qn('w:val'),'single');top.set(qn('w:sz'),'6');borders.append(top);p._p.get_or_add_pPr().append(borders)
    lines=[
        'Input: source S, operation o, recipient r, policy P',
        'Output: checked result, or explicit failure',
        '1:  if LocalRule(o, S) exists then return LocalRule(o, S)',
        '2:  K ← InitialFields(o, r, S); rounds ← 0',
        '3:  while rounds < 6 do',
        '4:      V ← Project(S, K); B ← Serialize(o, V)',
        '5:      T ← Bind(o, r, P.version, S[K], B)',
        '6:      if not CurrentAndUnchanged(T) then fail',
        '7:      y, receipt ← SendOnce(T, B); Verify(receipt, B)',
        '8:      if y requests fields G then',
        '9:          RequireAllowedNewFields(G, o, r, S, K)',
        '10:         K ← K ∪ G; EnforceAcquisitionLimit(2)',
        '11:     else if y requests a registered reference then',
        '12:         S ← Merge(S, CheckedLookup(y))',
        '13:         K ← RebuildFields(o, r, S, K)',
        '14:     else return ValidateAndAssembleLocally(y, S)',
        '15:     rounds ← rounds + 1',
        '16: fail STEP_LIMIT',
    ]
    for i,line in enumerate(lines):
        q=doc.add_paragraph();q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=1.05;q.paragraph_format.keep_with_next=i<len(lines)-1
        run(q,line,8.3,bold=i<2)
        if i==len(lines)-1:
            b=OxmlElement('w:pBdr');n=OxmlElement('w:bottom');n.set(qn('w:val'),'single');n.set(qn('w:sz'),'6');b.append(n);q._p.get_or_add_pPr().append(b)


def table(doc,title,rows,widths):
    p=heading(doc,title,2)
    t=doc.add_table(rows=len(rows),cols=len(widths));t.autofit=False
    for j,w in enumerate(widths):
        t.columns[j].width=Inches(w);t._tbl.tblGrid.gridCol_lst[j].set(qn('w:w'),str(round(w*1440)))
    for i,row in enumerate(rows):
        for j,text in enumerate(row):
            c=t.cell(i,j);c.width=Inches(widths[j]);p=c.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1
            run(p,text,7.7,bold=i==0)
            borders=OxmlElement('w:tcBorders')
            for edge in (['top','bottom'] if i==0 else ['bottom'] if i==len(rows)-1 else []):
                n=OxmlElement('w:'+edge);n.set(qn('w:val'),'single');n.set(qn('w:sz'),'4');borders.append(n)
            c._tc.get_or_add_tcPr().append(borders)
