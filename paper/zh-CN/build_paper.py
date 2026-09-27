"""Four-page Chinese editorial draft with native editable OMML mathematics."""
from layout import *
from docx.enum.text import WD_BREAK

def column_break(doc):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.line_spacing=Pt(1);p.add_run().add_break(WD_BREAK.COLUMN)

def algorithm(doc):
    p=heading(doc,'算法 1  面向接收方的后续规划与局部修复')
    b=OxmlElement('w:pBdr');n=OxmlElement('w:top');n.set(qn('w:val'),'single');n.set(qn('w:sz'),'6');b.append(n);p._p.get_or_add_pPr().append(b)
    lines=['Input: versioned state S, graph Q, registry R, slack ρ', 'Output: checked results C and disclosure history H', '1:  C ← ∅; H ← ∅; B ← UNSET', '2:  while Q is not empty do', '3:      if StateChanged(S) then', '4:          C ← KeepValidPureResults(C, S)', '5:          Q ← Pending(Q, C); B ← UNSET', '6:      A ← EquivalentAlternatives(S, Q, R)', '7:      if B = UNSET then', '8:          B ← floor((1 + ρ) × Fields(Greedy(A,H)))', '9:      p ← BudgetFrontier(A, H, B)', '10:     if p is infeasible then return BLOCKED', '11:     a ← First(p); CheckCurrent(S, a)', '12:     if a.recipient = local then y ← Evaluate(a)', '13:     else', '14:         bytes ← Serialize(a); CheckCurrent(S, a)', '15:         y, receipt ← SendOnce(bytes)', '16:         VerifyReceipt(receipt, bytes)', '17:         H ← H ∪ DisclosureTokens(a)', '18:     B ← B − NumericFields(a)', '19:     C ← SavePureResult(C, y, SourceDeps(a))', '20:     Q ← RemoveCompleted(Q, a)', '21: return C, H']
    for i,line in enumerate(lines):
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.02;p.paragraph_format.keep_with_next=i<len(lines)-1
        run(p,line,8,bold=i<2)
        if i==len(lines)-1:
            b=OxmlElement('w:pBdr');n=OxmlElement('w:bottom');n.set(qn('w:val'),'single');n.set(qn('w:sz'),'6');b.append(n);p._p.get_or_add_pPr().append(b)

def build():
    doc=Document();page(doc.sections[0],1)
    doc.core_properties.title='面向动态智能体任务的按步执行与信息共享'
    doc.core_properties.author='Yaojunzhe Zhai; Chao-Hsien Hsieh'
    doc.core_properties.subject='系统演示：接收方披露历史、后续规划与局部修复'
    doc.core_properties.comments=''
    normal=doc.styles['Normal'];normal.font.name='Linux Libertine';normal.font.size=Pt(9);normal.paragraph_format.space_after=Pt(0)
    p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3)
    b=OxmlElement('w:pBdr');n=OxmlElement('w:bottom');n.set(qn('w:val'),'nil');b.append(n);p._p.get_or_add_pPr().append(b)
    run(p,'面向动态智能体任务的按步执行与信息共享',14,True,cjk='SimHei')
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(2)
    run(p,'Yaojunzhe Zhai¹  ·  Chao-Hsien Hsieh²',9)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4)
    run(p,'¹ Monash University Malaysia   ² Xi’an International University',8)
    heading(doc,'摘要')
    paragraph(doc,'智能体运行中出现的新步骤，可能需要不同的执行位置与输入；源值或能力变化又会使已有计划失效。本文演示按步执行与信息共享系统：结合接收方已见信息选择当前已知步骤的等价输入，在数值字段传输预算内规划后续执行，并保留仍有效的纯计算结果。24 道公开财务问题的 96 次真实 DeepSeek 调用产生 48 份计划；在五种受控条件、六方法的配对执行中，局部修复比完整重跑减少 21.1% 的 HTTP 算子调用。后续规划在该批次未优于逐步贪心；48 份计划均可执行，但仅 28 份与原始数值标签一致。双栏界面将模型提出的步骤、条件变化与实际接收内容对应展示，使执行收益及其边界可以核对。',first=False,size=9)
    paragraph(doc,'关键词：智能体执行；接收方视图；累计披露；局部修复；系统演示。',first=False,size=8)
    heading(doc,'1  问题与系统目标')
    paragraph(doc,'以合同计费任务为例：一个步骤计算金额，另一个步骤检查限额。外部服务可以接收原始数值，也可以接收本地计算的金额。只看当前步骤，发送一个派生值似乎更省；考虑后续步骤，向同一接收方复用已披露的输入可能更合适。若限额在发送前改变，旧请求应当失效，但不依赖限额的已完成计算仍可保留。')
    full_figure(doc,'assets/system.png','图 1  按步执行的系统结构。原始资料保留在本地；模型与工具通过登记出口交互。新步骤或状态变化触发重新判定，外部结果写回本地状态。')
    col_section(doc)
    paragraph(doc,'这涉及三个相互关联的数据管理问题：计算放在哪里、每个接收方累计获知什么、状态变化后哪些结果仍有效。图 1 展示执行边界；本文在这一结构上实现可追溯的后续规划与局部修复。')
    paragraph(doc,'PlanTwin [1] 已提出云端规划抽象与累计披露预算；ToolMinimize [2] 和 MINIM [3] 分别限制工具参数与界面观察。ATR [4] 研究变化后的决策重验证，FreshCtx [5] 提供上下文时效检查。因此，裁剪、预算和依赖核验本身均不足以构成新颖性。')
    column_break(doc)
    paragraph(doc,'本文将执行位置、不可撤回的披露历史与变化后的结果复用联合处理。具体实现比较当前已知步骤的登记等价输入，在一个阶段共用剩余传输预算，并用来源依赖区分需要重算和可以保留的结果。贡献是这一受限执行机制及可核验演示，而非新的通用动态规划或自动语义脱敏算法。')
    paragraph(doc,'演示让观众回答三个具体问题：模型提出了哪些计算；条件变化后哪些步骤仍有效；外部接收方究竟收到哪些数值。预算约束限制重复发送字段，已披露集合则记录不同信息项，两者不能混为一谈。真实模型计划和受控干预保存在同一条执行链中，错误答案与无收益结果也予以保留。')

    col_section(doc,WD_SECTION_START.NEW_PAGE,2)
    heading(doc,'2  面向接收方的后续执行')
    heading(doc,'2.1  执行候选与披露历史',2)
    paragraph(doc,'版本化状态 S 保存源值、授权与本地能力。模型先申请不含业务数值的表结构，再提交有限计算图；控制器仅接受六种登记算子、前向执行且无无用节点的程序。候选 a 指定接收方 r 与等价视图 V：原始字段加子表达式，或已计算的直接操作数。')
    paragraph(doc,'一次披露以接收方、源字段或派生表达式及来源版本的指纹标识。令 E(a) 为候选披露项，H 为确认发送的集合；本地执行的 E(a) 为空。每次得到匹配接收记录后更新：')
    math(doc,'<msub><mi>H</mi><mrow><mi>t</mi><mo>+</mo><mn>1</mn></mrow></msub><mo>=</mo><msub><mi>H</mi><mi>t</mi></msub><mo>∪</mo><mi>E</mi><mo>(</mo><msub><mi>a</mi><mi>t</mi></msub><mo>)</mo>',1)
    paragraph(doc,'同一字段发送给不同接收方分别计数；派生值也计数。源版本变化后再次发送计为新项。H 不随重试或局部修复回退。该指标描述显式披露，不度量敏感性，也不排除接收方通过多个值联合推断。')
    heading(doc,'2.2  比较当前已知的后续方案',2)
    paragraph(doc,'设 Q 为当前已知待执行图，P(Q,S) 为满足授权、能力和剩余预算的候选组合；W(p) 是本地及接收端实际求值的算子数，T(p) 是发送的数值字段次数。原始字段重复发送仍计入 T，新增披露集合为：')
    math(doc,'<mi>D</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>=</mo><mo>(</mo><munder><mo>⋃</mo><mrow><mi>a</mi><mo>∈</mo><mi>p</mi></mrow></munder><mi>E</mi><mo>(</mo><mi>a</mi><mo>)</mo><mo>)</mo><mo>∖</mo><mi>H</mi>',2)
    math(doc,'<msup><mi>p</mi><mo>∗</mo></msup><mo>=</mo><munder><mtext>arg min lex</mtext><mrow><mi>p</mi><mo>∈</mo><mi>P</mi><mo>(</mo><mi>Q</mi><mo>,</mo><mi>S</mi><mo>)</mo></mrow></munder><mo>(</mo><mo>|</mo><mi>D</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>|</mo><mo>,</mo><mi>W</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>,</mo><mi>T</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>)</mo>',3)
    paragraph(doc,'式（3）依次比较不同披露项、算子求值与传输字段。阶段开始时，以相同状态下逐步贪心的传输量为基准设预算 B；默认不增加字段，敏感性设置允许增加 10% 或 25%。执行每一步都扣减同一余额，不能重新取得整份额度。')
    math(doc,'<mi>B</mi><mo>=</mo><mo>⌊</mo><mo>(</mo><mn>1</mn><mo>+</mo><mi>ρ</mi><mo>)</mo><mi>T</mi><mo>(</mo><msub><mi>p</mi><mi>g</mi></msub><mo>)</mo><mo>⌋</mo>',4)
    paragraph(doc,'同一深度的标签记录披露集合、累计求值与传输 (E,w,b)。下式成立时第二个标签可被剪枝：')
    math(doc,'<msub><mi>E</mi><mn>1</mn></msub><mo>⊆</mo><msub><mi>E</mi><mn>2</mn></msub><mo>∧</mo><msub><mi>w</mi><mn>1</mn></msub><mo>≤</mo><msub><mi>w</mi><mn>2</mn></msub><mo>∧</mo><msub><mi>b</mi><mn>1</mn></msub><mo>≤</mo><msub><mi>b</mi><mn>2</mn></msub>',5)
    paragraph(doc,'若候选输出等价且后续可行性不依赖表示，追加同一后缀保持集合包含与两项成本关系，因此剪枝不会丢失受限问题的最优解。未来派生值用表达式及源版本标识，无需预先计算其数值。最坏搜索仍为指数级，超过 4,096 个标签明确失败。')
    heading(doc,'一个逐步选择失效的例子',2)
    paragraph(doc,'设两个步骤均可向服务甲发送同一对原值 {x,y}；另一条路径向服务乙分别发送一个派生值 {u} 与两个不同派生值 {v,w}。逐步贪心先选乙，只披露一项；第二步无论选哪条路径，累计都至少为三项。后续规划连续选择甲，累计只有两项，但会传输四个字段。该例说明为什么披露集合与传输次数须分别衡量。')
    column_break(doc)
    heading(doc,'2.3  核验决策，而不只核验请求字段',2)
    paragraph(doc,'请求内容未变，不代表原执行方式仍可行。本地能力变化可能要求改用外部算子；源值变化可能使缓存过期。控制器区分决定候选可行性的规划依赖与决定结果数值的传递依赖：')
    math(doc,'<mi>Δ</mi><mo>=</mo><msub><mi>Δ</mi><mtext>值</mtext></msub><mo>∪</mo><msub><mi>Δ</mi><mtext>规划</mtext></msub><mo>,</mo><mspace width="0.3em"/><mo>∀</mo><mi>k</mi><mo>∈</mo><mi>Δ</mi><mo>:</mo><mi>v</mi><mo>(</mo><mi>k</mi><mo>)</mo><mo>=</mo><msub><mi>v</mi><mtext>当前</mtext></msub><mo>(</mo><mi>k</mi><mo>)</mo>',6)
    paragraph(doc,'源版本、当前能力和授权在序列化前同步检查；状态变化后重新选择剩余步骤，缓存仅按值依赖判断是否保留。受控版本往返测试使用单调版本号。已失效结果不得作为新请求的操作数，未用源字段改变则不使该结果失效。')
    algorithm(doc)
    paragraph(doc,'算法 1 在每个状态稳定阶段维持一个余额。阶段内，剩余旧方案为下一次受限选择提供可行解；累计实际发送字段因而不超过阶段预算。新的干预会开启新阶段，这个局部界限不能替代整轮任务的实测传输量。')
    paragraph(doc,'接收端记录原始字节、摘要与结果，核验器逐请求重算。新模型图实验使用同一 Node 进程内的回环 HTTP 监听器；真实模型请求由另一个代理保存供应商载荷与响应摘要。它们均不构成供应商签名回执或生产分布式事务证明。')
    paragraph(doc,'自动复用仅限登记的无副作用计算。接收授权撤销不能收回既往披露；网络返回不确定时停止，不假设未获回执就没有信息外发。任意脚本、带副作用工具及不可信控制器均不在保证范围内。')
    paragraph(doc,'该预算只约束当前状态与已知后续图。新步骤或源状态变化后必须重建阶段，不能由此推导跨变化的总量界限。字段数也不是 HTTP 字节数；报文标签、表达式与结构信息另占空间，实验分别记录这些代价。')

    col_section(doc,WD_SECTION_START.NEW_PAGE,1)
    full_figure(doc,'../../evidence/model-showcase-v8/paper-ui.png','图 2  公开财务表格的实际执行界面。① 从保存的真实模型回复创建计算步骤；② 受控撤销本地能力，保留已完成结果；③ HTTP 接收端收到中间值与所需源值。截图为冻结批次中一个计划的独立重执行，裁去导航并添加编号，不增加测试样本数。',width=6.65)
    col_section(doc)
    heading(doc,'3  交互演示')
    paragraph(doc,'界面沿用客户端的双栏任务布局。左侧执行自动推进，右侧随选中步骤展示接收方、发送内容、保留字段与接收记录。观众可暂停跟随检查旧请求，再回到当前步骤；无需逐步点击才能完成任务。展示节奏单独放慢，不计入实验耗时。')
    paragraph(doc,'第一组使用公开财务表格：模型通过真实工具调用申请表结构，随后提交“相加、相加、相除”的计算图。观众查看原始回复后执行该图；首步完成时撤销登记的本地加减能力，系统保留首步结果，将后续步骤交给 HTTP 算子。切换完整重跑，可直接看到该首步被重复执行。')
    paragraph(doc,'第二组分别改变已用与未用源值。已用值更新使依赖它的结果失效；未用值更新则保留已有计算。右侧显示原始字段或中间值，以及本阶段剩余字段预算。全本地条件展示无需外部数值服务的边界，混合能力是实验配置，不表示加减乘除天然需要云端。')
    paragraph(doc,'无密钥模式重用公开的真实模型计划，但重新执行本地计算及 HTTP 请求；不是重新调用模型，也不是只播放预设结果。读者可用自己的密钥运行完整批次命令，产生新的模型计划。旧版合同案例、逐步补充和人工撤权记录继续保留，均与新实验分开。')
    column_break(doc)
    heading(doc,'4  实验设计')
    paragraph(doc,'在新调用前冻结提示词、执行器、评价器和事件安排。从 FinQA [7] 排除此前使用的公司年度报告，固定顺序选择 24 题、24 份报告；12 题来自原 test 划分、12 题来自原 dev 划分，各保留来源标记。兼容筛选限定可由表格与登记算子回答的多操作问题；另四题仅用于开发。')
    paragraph(doc,'每题实际调用 DeepSeek 两轮生成计划，每份计划包含申请结构与提交计算两个请求，共 48 份计划、96 次调用。每份计划在五种条件下比较六方法：完整重跑、修复加贪心、无传输约束的后续规划，以及 0%、10%、25% 预算。共享模型前缀保证执行方法的配对可比；1,440 次执行不是 1,440 次独立模型任务。')
    paragraph(doc,'五种条件为全本地、混合能力无变化、已用值更新、未用值更新和本地能力撤销。干预在模型图约半数步骤完成时触发，源值按预定规则变化；执行器不读参考答案。评估分别检查模型图执行一致性、原始标签一致性、来源有效性及实际载荷，全部错误保留。')
    paragraph(doc,'另保存 288 个合成变化场景、六方法的机制消融，以及旧版 600 次模型任务；它们不与本批次合并统计。新批次每道题的两份计划和五种条件相关，统计单位为题。按题聚类的探索性自助法只用于描述外部调用减少量，不证明真实生产中的泛化收益。')

    col_section(doc,WD_SECTION_START.NEW_PAGE,1)
    full_figure(doc,'../../evidence/validation-v8/tradeoffs.zh-CN.png','图 3  新模型图实验，每种方法 240 次配对执行。a、b：实际 HTTP 调用与数值字段传输；c：按干预条件拆分；d：R、G、F 分别为重跑、贪心与后续规划，后三项为传输预算设置。本批次的五种修复方法结果相同，不支持额外规划收益。',width=6.78)
    col_section(doc)
    heading(doc,'5  结果与讨论')
    paragraph(doc,'48 份模型图均可执行，各方法在 240 次运行中均与该图的当前数值语义一致，但只有 137 次与相应的参考表达式输出一致。原始、未干预问题的标签一致性为 28/48。这里“执行正确”只说明按模型表达式算对，不能说明模型选对了字段或理解了题意。')
    paragraph(doc,'与完整重跑相比，保留有效结果将外部调用从 331 降到 261，减少 21.1%，复用 112 个计算；传输数值字段由 724 降至 500。按 24 题聚类、10,000 次探索性重采样，调用减少比例的区间为 [19.1%, 23.8%]。主要收益来自能力撤销及未用值变化后的结果保留。')
    paragraph(doc,'逐步贪心、后续规划及三档预算均产生 425 个不同披露项、500 个传输字段，未观察到后续规划的新增收益。旧合成测试中规划曾减少 2.9% 的披露项，却增加 54.5% 的传输字段；新约束提供阶段预算界限，但本批次并未证明该约束带来额外效率提升。')
    paragraph(doc,'错误记录包括年份选择、差值方向和百分比单位；另两道平均数题的原始标签程序疑似与题意不符，仍保留原评分并公开复核材料。未作独立人审前，标签一致率不能当作已确认的语义正确率。报文字节另受运行标识长度影响，不能把命名开销差异解释成算法收益。')
    column_break(doc)
    heading(doc,'6  适用范围与开放实现')
    paragraph(doc,'当前证据支持真实模型计划进入受限执行闭环，以及变化后复用有效结果的系统价值。预算是对已知图的可检验约束；后续规划相对强贪心的实际收益尚未建立。自动生成等价视图、更多任务域和模型、真人质量评阅及生产部署仍需验证。')
    paragraph(doc,'仓库提供独立源码、双语界面、无密钥运行、原始模型响应、逐请求接收记录及离线核验器。模型程序与披露决策可追溯到同一轮数据；旧实验不被覆盖。该演示让观众检查每一步为何执行、发送了什么，以及条件变化后哪些工作可以保留。',space=1)
    paragraph(doc,'代码与复现：https://github.com/Zane-0260907/stepwise-disclosure-demo',first=False,size=7.7)
    heading(doc,'参考文献')
    refs=[
      '[1] G. Yu 等. PlanTwin: Privacy-Preserving Planning Abstractions for Cloud-Assisted LLM Agents. arXiv:2603.18377, 2026.',
      '[2] W. Li, Y. Xu. ToolMinimize: Auditing and Rewriting LLM Agent Tool Calls to Minimize Privacy Exposure. arXiv:2608.24957, 2026.',
      '[3] H. Yu 等. MINIM: Privacy-Aware Minimal View for Agents via Trusted Local Sanitization. arXiv:2606.13949, 2026.',
      '[4] Y. Lyu, Y. Ren, R. Lai, W. Liu. From Version Conflicts to Decision Conflicts: Selective Revalidation for Long-Running AI Agents. arXiv:2609.08015, 2026.',
      '[5] Hyperwise. FreshCtx 0.16.0. github.com/Hyperwise-LLC/freshctx, 2026. Apache-2.0.',
      '[6] L. Gao 等. PAL: Program-aided Language Models. ICML, 2023, 10764–10799.',
      '[7] Z. Chen 等. FinQA: A Dataset of Numerical Reasoning over Financial Data. EMNLP, 2021, 3697–3711.'
    ]
    for text in refs:
        p=paragraph(doc,text,first=False,size=8,space=1,align=WD_ALIGN_PARAGRAPH.LEFT);p.paragraph_format.left_indent=Inches(.14);p.paragraph_format.first_line_indent=Inches(-.14)
    view=OxmlElement('w:view');view.set(qn('w:val'),'print');doc.settings.element.insert(0,view)
    doc.save(OUT);print(OUT)

if __name__=='__main__':build()
